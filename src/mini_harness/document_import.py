"""把本地文档转换成知识库标准 Markdown 的可插拔工具。"""

from __future__ import annotations

import csv
import io
import json
import re
from collections.abc import Callable, Mapping, Sequence
from dataclasses import asdict, dataclass
from datetime import date
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from .client import Tool
from .cordis import Context, depends
from .knowledge_base import KnowledgeBaseService
from .prompt import PromptAssembler
from .registry import ToolRegistry
from .settings import SettingsScope

SUPPORTED_SUFFIXES = frozenset(
    {".pdf", ".docx", ".md", ".txt", ".html", ".htm", ".json", ".csv"}
)
H1 = re.compile(r"^\s*#\s+(.+?)\s*$", re.MULTILINE)


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    title: str | None
    source_format: str
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class ImportedDocument:
    source: str
    destination: str
    title: str
    source_format: str
    characters: int
    warnings: tuple[str, ...]

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False, allow_nan=False)


class DocumentImportService:
    """读取 workspace 内文档，并通过文件服务写入知识库。"""

    def __init__(
        self,
        *,
        workspace: Path,
        knowledge_base: KnowledgeBaseService,
        writer: Callable[[dict[str, Any]], str],
        import_dir: str,
        max_bytes: int,
    ) -> None:
        self.workspace = workspace.resolve()
        self.knowledge_base = knowledge_base
        self.writer = writer
        self.import_dir = import_dir
        self.max_bytes = max_bytes

    def import_document(
        self,
        *,
        source_path: str,
        destination: str | None = None,
        title: str | None = None,
        tags: Sequence[str] = (),
    ) -> ImportedDocument:
        source = self._source_path(source_path)
        suffix = source.suffix.casefold()
        if suffix not in SUPPORTED_SUFFIXES:
            supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
            raise ValueError(f"不支持 {suffix or '无扩展名'}；支持格式: {supported}")
        size = source.stat().st_size
        if size > self.max_bytes:
            raise ValueError(
                f"源文件为 {size} bytes，超过上限 {self.max_bytes} bytes"
            )

        extracted = _extract(source)
        body = extracted.text.strip()
        if not body:
            raise ValueError("源文件没有提取出可导入文本")
        candidate_title = title or extracted.title or _title_from(body)
        resolved_title = _clean_inline(candidate_title) if candidate_title else ""
        if not resolved_title:
            resolved_title = _clean_inline(source.stem.replace("_", " ").replace("-", " "))
        if not resolved_title:
            raise ValueError("无法推断文档标题，请显式传入 title")

        target = self._destination(source, destination)
        if target.exists():
            raise FileExistsError(f"目标笔记已存在，拒绝覆盖: {target}")
        clean_tags = tuple(
            dict.fromkeys(_clean_inline(tag) for tag in tags if _clean_inline(tag))
        ) or ("imported",)
        relative_source = source.relative_to(self.workspace).as_posix()
        markdown = _render_markdown(
            title=resolved_title,
            tags=clean_tags,
            source=relative_source,
            body=body,
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        self.writer({"path": str(target), "content": markdown})
        return ImportedDocument(
            source=relative_source,
            destination=target.relative_to(self.knowledge_base.root).as_posix(),
            title=resolved_title,
            source_format=extracted.source_format,
            characters=len(body),
            warnings=extracted.warnings,
        )

    def _source_path(self, value: str) -> Path:
        if not value.strip():
            raise ValueError("source_path 必须是非空路径")
        raw = Path(value)
        source = raw.resolve() if raw.is_absolute() else (self.workspace / raw).resolve()
        try:
            source.relative_to(self.workspace)
        except ValueError as error:
            raise PermissionError("source_path 必须位于当前 workspace 内") from error
        if not source.is_file():
            raise FileNotFoundError(f"源文件不存在: {source}")
        return source

    def _destination(self, source: Path, value: str | None) -> Path:
        relative = (
            Path(value)
            if value
            else Path(self.import_dir) / f"{_safe_stem(source.stem)}.md"
        )
        if relative.is_absolute() or relative.suffix.casefold() != ".md":
            raise ValueError("destination 必须是知识库根目录内的相对 .md 路径")
        target = (self.knowledge_base.root / relative).resolve()
        try:
            target.relative_to(self.knowledge_base.root)
        except ValueError as error:
            raise PermissionError("destination 不能逃出知识库根目录") from error
        return target


@depends("knowledge_base", "filesystem", "agent_settings")
def document_import_provider(ctx: Context, _config: Any) -> None:
    knowledge_base = ctx.knowledge_base
    filesystem = ctx.filesystem
    settings = ctx.agent_settings
    assert isinstance(knowledge_base, KnowledgeBaseService)
    assert isinstance(settings, SettingsScope)
    config = settings.get()
    service = DocumentImportService(
        workspace=Path(filesystem.workspace),
        knowledge_base=knowledge_base,
        writer=filesystem.write,
        import_dir=_required_setting_str(config["kb_import_dir"], "kb_import_dir"),
        max_bytes=_positive_int(config["kb_import_max_bytes"], "kb_import_max_bytes"),
    )
    ctx.provide("document_import", service)


@depends("document_import", "tools", "prompt")
def document_import_consumer(ctx: Context, _config: Any) -> None:
    service = ctx.document_import
    tools = ctx.tools
    prompt = ctx.prompt
    assert isinstance(service, DocumentImportService)
    assert isinstance(tools, ToolRegistry)
    assert isinstance(prompt, PromptAssembler)
    ctx.effect(
        lambda: prompt.section(
            "knowledge-base:import",
            (
                "导入用户指定的本地 PDF、DOCX 或文本类文件时使用 "
                "kb_import_document。源文件和目标必须位于 workspace/knowledge "
                "base 内；工具拒绝覆盖已有笔记。导入后用 read 核对结果，必要时"
                "再做最小编辑，最后可用 kb_audit 复审。"
            ),
            order=11,
        )
    )
    ctx.effect(
        lambda: tools.register(
            Tool(
                "kb_import_document",
                (
                    "把 workspace 内的 PDF、DOCX、Markdown、TXT、HTML、JSON "
                    "或 CSV 转换成带标准 frontmatter 的知识库 Markdown 笔记。"
                ),
                {
                    "type": "object",
                    "properties": {
                        "source_path": {"type": "string", "minLength": 1},
                        "destination": {
                            "type": "string",
                            "description": "知识库根目录内的相对 .md 路径",
                        },
                        "title": {"type": "string"},
                        "tags": {
                            "type": "array",
                            "items": {"type": "string"},
                            "uniqueItems": True,
                        },
                    },
                    "required": ["source_path"],
                },
                lambda arguments: _import_tool(service, arguments),
            )
        )
    )


def _import_tool(
    service: DocumentImportService, arguments: Mapping[str, Any]
) -> str:
    source_path = arguments.get("source_path")
    if not isinstance(source_path, str) or not source_path.strip():
        raise ValueError("source_path 必须是非空字符串")
    destination = arguments.get("destination")
    if destination is not None and not isinstance(destination, str):
        raise TypeError("destination 必须是字符串")
    title = arguments.get("title")
    if title is not None and not isinstance(title, str):
        raise TypeError("title 必须是字符串")
    raw_tags = arguments.get("tags", [])
    if not isinstance(raw_tags, list) or not all(
        isinstance(tag, str) for tag in raw_tags
    ):
        raise TypeError("tags 必须是字符串数组")
    return service.import_document(
        source_path=source_path,
        destination=destination,
        title=title,
        tags=raw_tags,
    ).to_json()


def _extract(path: Path) -> ExtractedDocument:
    suffix = path.suffix.casefold()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".md":
        text = path.read_text(encoding="utf-8")
        fields, body = _split_frontmatter(text)
        return ExtractedDocument(body, fields.get("title"), "markdown")
    if suffix == ".txt":
        return ExtractedDocument(path.read_text(encoding="utf-8"), None, "text")
    if suffix in {".html", ".htm"}:
        parser = _TextHtmlParser()
        parser.feed(path.read_text(encoding="utf-8"))
        return ExtractedDocument(parser.text(), parser.title, "html")
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        pretty = json.dumps(value, ensure_ascii=False, indent=2)
        return ExtractedDocument(f"```json\n{pretty}\n```", None, "json")
    if suffix == ".csv":
        raw = path.read_text(encoding="utf-8-sig")
        # 先完整解析一次，确保损坏的 CSV 在写入前失败。
        list(csv.reader(io.StringIO(raw)))
        return ExtractedDocument(f"```csv\n{raw.rstrip()}\n```", None, "csv")
    raise ValueError(f"不支持的文档格式: {suffix}")


def _extract_pdf(path: Path) -> ExtractedDocument:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as error:  # noqa: BLE001 - 第三方解析边界
            raise ValueError("PDF 已加密且无法读取") from error
        if not unlocked:
            raise ValueError("PDF 已加密且需要密码")
    pages: list[str] = []
    empty_pages = 0
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            empty_pages += 1
            continue
        pages.append(f"## 第 {index} 页\n\n{text}")
    metadata_title = None
    if reader.metadata is not None:
        metadata_title = getattr(reader.metadata, "title", None)
    if not pages and empty_pages:
        raise ValueError("PDF 未提取到文本；扫描版 PDF 需要 OCR（当前未支持）")
    warnings = (
        (f"{empty_pages} 页未提取到文本，扫描版 PDF 需要 OCR",)
        if empty_pages
        else ()
    )
    return ExtractedDocument(
        "\n\n".join(pages), metadata_title, "pdf", warnings
    )


def _extract_docx(path: Path) -> ExtractedDocument:
    document = Document(str(path))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        heading = re.fullmatch(r"Heading (\d+)", style)
        if heading:
            level = min(max(int(heading.group(1)), 1), 6)
            blocks.append(f"{'#' * level} {text}")
        else:
            blocks.append(text)
    for table in document.tables:
        rows = [
            [_escape_table(cell.text.strip()) for cell in row.cells]
            for row in table.rows
        ]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        blocks.append("| " + " | ".join(normalized[0]) + " |")
        blocks.append("| " + " | ".join("---" for _ in range(width)) + " |")
        blocks.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    title = document.core_properties.title or None
    return ExtractedDocument("\n\n".join(blocks), title, "docx")


class _TextHtmlParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []
        self._in_title = False
        self._title_parts: list[str] = []
        self._ignored_depth = 0
        self.title: str | None = None

    def handle_starttag(
        self, tag: str, _attrs: list[tuple[str, str | None]]
    ) -> None:
        if tag == "title":
            self._in_title = True
        if tag in {"script", "style"}:
            self._ignored_depth += 1
        if tag in {"p", "div", "br", "li", "h1", "h2", "h3", "tr"}:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self._in_title = False
            value = _clean_inline("".join(self._title_parts))
            self.title = value or None
        if tag in {"script", "style"} and self._ignored_depth:
            self._ignored_depth -= 1
        if tag in {"p", "div", "li", "h1", "h2", "h3", "tr"}:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self._title_parts.append(data)
        elif not self._ignored_depth:
            self._parts.append(data)

    def text(self) -> str:
        lines = (" ".join(line.split()) for line in "".join(self._parts).splitlines())
        return "\n\n".join(line for line in lines if line)


def _render_markdown(
    *, title: str, tags: Sequence[str], source: str, body: str
) -> str:
    metadata = [
        "---",
        f"title: {_frontmatter_value(title)}",
        f"tags: {_frontmatter_value(', '.join(tags))}",
        f"source: {_frontmatter_value(source)}",
        f"updated: {date.today().isoformat()}",
        "---",
        "",
    ]
    if not H1.search(body):
        metadata.extend((f"# {title}", ""))
    metadata.extend((body, ""))
    return "\n".join(metadata)


def _split_frontmatter(text: str) -> tuple[dict[str, str], str]:
    lines = text.splitlines()
    if not lines or lines[0].strip() != "---":
        return {}, text
    fields: dict[str, str] = {}
    for index, line in enumerate(lines[1:], start=1):
        if line.strip() == "---":
            return fields, "\n".join(lines[index + 1 :])
        if ":" in line:
            key, _, value = line.partition(":")
            fields[key.strip()] = value.strip().strip("\"'")
    return {}, text


def _title_from(text: str) -> str | None:
    match = H1.search(text)
    return _clean_inline(match.group(1)) if match else None


def _frontmatter_value(value: str) -> str:
    return json.dumps(_clean_inline(value), ensure_ascii=False)


def _clean_inline(value: str) -> str:
    return " ".join(value.replace("\x00", "").split())


def _safe_stem(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "-", value).strip(" .-")
    return cleaned or "imported-document"


def _escape_table(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>")


def _required_setting_str(value: Any, name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise TypeError(f"{name} 必须是非空字符串")
    return value


def _positive_int(value: Any, name: str) -> int:
    if not isinstance(value, int) or isinstance(value, bool) or value <= 0:
        raise TypeError(f"{name} 必须是正整数")
    return value
