from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

from mini_harness.agent import Agent
from mini_harness.bundle import BundleConfig, headless_bundle
from mini_harness.client import Message, Tool
from mini_harness.cordis import Context
from mini_harness.document_import import DocumentImportService
from mini_harness.knowledge_base import KnowledgeBaseService
from mini_harness.prompt import PromptAssembler
from mini_harness.registry import ToolRegistry


class FakeClient:
    MODEL = "fake"

    def chat(
        self, _messages: list[Message], _tools: list[Tool] | None = None
    ) -> Message:
        return Message(role="assistant", content="unused")


class DocumentImportServiceTests(unittest.TestCase):
    def test_text_is_normalized_to_standard_markdown(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            _write(workspace / "inbox" / "agent.txt", "插件生命周期与依赖装配。")
            service = _service(workspace)

            result = service.import_document(
                source_path="inbox/agent.txt",
                title="Agent 架构",
                tags=["Python", "Harness", "Python"],
            )

            target = workspace / "kb" / "imports" / "agent.md"
            content = target.read_text(encoding="utf-8")
            self.assertEqual(result.destination, "imports/agent.md")
            self.assertEqual(result.source_format, "text")
            self.assertIn('title: "Agent 架构"', content)
            self.assertIn('tags: "Python, Harness"', content)
            self.assertIn('source: "inbox/agent.txt"', content)
            self.assertIn("# Agent 架构", content)
            self.assertIn("插件生命周期与依赖装配。", content)

    def test_docx_extracts_headings_paragraphs_tables_and_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            source = workspace / "inbox" / "design.docx"
            source.parent.mkdir(parents=True)
            document = Document()
            document.core_properties.title = "插件设计"
            document.add_heading("核心架构", level=1)
            document.add_paragraph("内核只负责生命周期。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "能力"
            table.cell(0, 1).text = "实现"
            table.cell(1, 0).text = "工具"
            table.cell(1, 1).text = "Consumer"
            document.save(source)

            result = _service(workspace).import_document(source_path="inbox/design.docx")

            content = (workspace / "kb" / "imports" / "design.md").read_text(
                encoding="utf-8"
            )
            self.assertEqual(result.title, "插件设计")
            self.assertIn("# 核心架构", content)
            self.assertIn("内核只负责生命周期。", content)
            self.assertIn("| 能力 | 实现 |", content)
            self.assertIn("| 工具 | Consumer |", content)

    def test_markdown_frontmatter_is_replaced_instead_of_nested(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            _write(
                workspace / "source.md",
                "---\ntitle: 原始标题\nlegacy: yes\n---\n\n# 正文标题\n\n正文。",
            )

            result = _service(workspace).import_document(source_path="source.md")

            content = (workspace / "kb" / "imports" / "source.md").read_text(
                encoding="utf-8"
            )
            self.assertEqual(result.title, "原始标题")
            self.assertEqual(content.count("---"), 2)
            self.assertNotIn("legacy:", content)
            self.assertEqual(content.count("# 正文标题"), 1)

    def test_pdf_preserves_pages_and_reports_empty_pages(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            source = workspace / "report.pdf"
            source.write_bytes(b"%PDF-test")
            reader = SimpleNamespace(
                is_encrypted=False,
                metadata=SimpleNamespace(title="季度报告"),
                pages=[_PdfPage("第一页"), _PdfPage(""), _PdfPage("第三页")],
            )

            with patch("mini_harness.document_import.PdfReader", return_value=reader):
                result = _service(workspace).import_document(source_path="report.pdf")

            content = (workspace / "kb" / "imports" / "report.md").read_text(
                encoding="utf-8"
            )
            self.assertEqual(result.title, "季度报告")
            self.assertEqual(result.warnings, ("1 页未提取到文本，扫描版 PDF 需要 OCR",))
            self.assertIn("## 第 1 页", content)
            self.assertIn("## 第 3 页", content)

    def test_scanned_or_blank_pdf_requires_ocr(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            source = workspace / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with source.open("wb") as stream:
                writer.write(stream)

            with self.assertRaisesRegex(ValueError, "OCR"):
                _service(workspace).import_document(source_path="blank.pdf")

    def test_html_json_and_csv_are_supported(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            _write(
                workspace / "page.html",
                "<title>页面</title><h1>正文</h1><p>内容</p><script>secret()</script>",
            )
            _write(workspace / "data.json", '{"name": "小知"}')
            _write(workspace / "data.csv", "name,value\n小知,1\n")
            service = _service(workspace)

            service.import_document(source_path="page.html")
            service.import_document(
                source_path="data.json", destination="imports/data-json.md"
            )
            service.import_document(
                source_path="data.csv", destination="imports/data-csv.md"
            )

            html = (workspace / "kb" / "imports" / "page.md").read_text(
                encoding="utf-8"
            )
            json_note = (workspace / "kb" / "imports" / "data-json.md").read_text(
                encoding="utf-8"
            )
            self.assertNotIn("secret()", html)
            self.assertIn('"name": "小知"', json_note)
            self.assertTrue((workspace / "kb" / "imports" / "data-csv.md").is_file())

    def test_rejects_root_escape_unsupported_format_and_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as other:
            workspace = Path(tmp)
            outside = Path(other) / "outside.txt"
            _write(outside, "private")
            _write(workspace / "source.txt", "content")
            _write(workspace / "source.doc", "legacy")
            service = _service(workspace)

            with self.assertRaises(PermissionError):
                service.import_document(source_path=str(outside))
            with self.assertRaises(PermissionError):
                service.import_document(
                    source_path="source.txt", destination="../escape.md"
                )
            with self.assertRaises(ValueError):
                service.import_document(source_path="source.doc")
            service.import_document(source_path="source.txt")
            with self.assertRaises(FileExistsError):
                service.import_document(source_path="source.txt")

    def test_rejects_files_over_configured_size_limit(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            _write(workspace / "large.txt", "12345")
            service = _service(workspace, max_bytes=4)

            with self.assertRaises(ValueError):
                service.import_document(source_path="large.txt")


class DocumentImportPluginTests(unittest.TestCase):
    def test_bundle_registers_tool_and_disposal_is_reversible(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            previous = Path.cwd()
            os.chdir(tmp)
            _write(Path("source.txt"), "需要导入的内容")
            ctx = Context()
            try:
                ctx.plugin(
                    headless_bundle,
                    BundleConfig(settings_document={}, client=FakeClient()),
                )
                agent = ctx.require("agent")
                tools = ctx.require("tools")
                prompt = ctx.require("prompt")
                self.assertIsInstance(agent, Agent)
                self.assertIsInstance(tools, ToolRegistry)
                self.assertIsInstance(prompt, PromptAssembler)
                self.assertEqual(len(agent.tool_names), 27)
                tool = tools.get("kb_import_document")
                self.assertIsNotNone(tool)
                assert tool is not None

                result = json.loads(tool.execute({"source_path": "source.txt"}))

                self.assertEqual(result["destination"], "source.md")
                self.assertTrue(Path("notes/source.md").is_file())
                self.assertIn(
                    "knowledge-base:import",
                    {section.name for section in prompt.sections},
                )

                provider = next(
                    handle
                    for handle in ctx.handles
                    if handle.name == "document_import_provider"
                )
                provider.dispose()

                self.assertNotIn("document_import", ctx.service_names)
                self.assertNotIn("kb_import_document", agent.tool_names)
                self.assertNotIn(
                    "knowledge-base:import",
                    {section.name for section in prompt.sections},
                )
            finally:
                ctx.dispose()
                os.chdir(previous)

    def test_disabled_knowledge_base_does_not_install_importer(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            previous = Path.cwd()
            os.chdir(tmp)
            ctx = Context()
            try:
                ctx.plugin(
                    headless_bundle,
                    BundleConfig(
                        settings_document={},
                        client=FakeClient(),
                        enable_knowledge_base=False,
                    ),
                )
                agent = ctx.require("agent")
                self.assertIsInstance(agent, Agent)
                self.assertEqual(len(agent.tool_names), 24)
                self.assertNotIn("kb_import_document", agent.tool_names)
                self.assertNotIn("document_import", ctx.service_names)
            finally:
                ctx.dispose()
                os.chdir(previous)


class _PdfPage:
    def __init__(self, text: str) -> None:
        self._text = text

    def extract_text(self) -> str:
        return self._text


def _service(workspace: Path, *, max_bytes: int = 1024 * 1024) -> DocumentImportService:
    root = workspace / "kb"
    knowledge_base = KnowledgeBaseService(
        root=root,
        include_globs=("**/*.md", "**/*.txt"),
        exclude_dirs=(".git", ".mini-harness", ".agents", "node_modules"),
        required_frontmatter=(),
        stale_days=180,
        index_pages=("README.md", "index.md", "INDEX.md"),
    )

    def writer(arguments: dict[str, object]) -> str:
        path = Path(str(arguments["path"]))
        path.write_text(str(arguments["content"]), encoding="utf-8")
        return "written"

    return DocumentImportService(
        workspace=workspace,
        knowledge_base=knowledge_base,
        writer=writer,
        import_dir="imports",
        max_bytes=max_bytes,
    )


def _write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


if __name__ == "__main__":
    unittest.main()
